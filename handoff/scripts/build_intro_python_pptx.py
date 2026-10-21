import sys, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
src,out=sys.argv[1],sys.argv[2]
p=Presentation(src)
NAVY=RGBColor(16,20,86);BLUE=RGBColor(33,62,219);INK=RGBColor(21,23,28);GRAY=RGBColor(97,103,115);LIGHT=RGBColor(238,241,252);WHITE=RGBColor(255,255,255)
F='Malgun Gothic'
def tb(s,x,y,w,h,text,size=14,color=INK,bold=False,fill=None,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Emu(x),Emu(y),Emu(w),Emu(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    else: sh.fill.background()
    if line: sh.line.color.rgb=line
    else: sh.line.fill.background()
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE if h<=450000 else MSO_ANCHOR.TOP
    for m in ('margin_left','margin_right'): setattr(tf,m,Emu(120000))
    lines=text if isinstance(text,list) else [text]
    for i,t in enumerate(lines):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if isinstance(t,tuple): t,sz,c,b=t
        else: sz,c,b=size,color,bold
        r=para.add_run(); r.text=t; r.font.name=F; r.font.size=Pt(sz); r.font.color.rgb=c; r.font.bold=b
        para.space_after=Pt(4); para.alignment=PP_ALIGN.LEFT
    return sh
def base(kicker,title,num,src_text):
    s=p.slides.add_slide(p.slide_layouts[6] if len(p.slide_layouts)>6 else p.slide_layouts[-1])
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    tb(s,590550,295275,9715500,228600,kicker,10.5,GRAY)
    tb(s,571500,742950,11049000,609600,title,28,INK,True)
    tb(s,590550,6200000,10500000,400000,src_text,9,GRAY)
    tb(s,11144250,6438900,457200,228600,num,10.5,NAVY,True)
    return s
def card(s,x,y,w,h,head,body,hfill=NAVY):
    tb(s,x,y,w,420000,head,14,WHITE,True,fill=hfill)
    tb(s,x,y+420000,w,h-420000,body,12.5,INK,fill=LIGHT)
def note(s,t): s.notes_slide.notes_text_frame.text=t
SRC1='출처: Anthropic, Handa·Tamkin 외 「Which Economic Tasks are Performed with AI? Evidence from Millions of Claude Conversations」(arXiv:2503.04761, 2025.2)'
SRC2='출처: Anthropic, Massenkoff·McCrory 「Labor market impacts of AI: A new measure and early evidence」(2026.3.5)'
X=[590550,4290000,7990000];W=3550000
new=[]
# 0 cover
s=p.slides.add_slide(p.slide_layouts[6] if len(p.slide_layouts)>6 else p.slide_layouts[-1])
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,p.slide_width,p.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb=NAVY; bg.line.fill.background()
tb(s,900000,2000000,10000000,500000,'INTRO  ·  도입',14,RGBColor(170,180,240),True)
tb(s,900000,2550000,10400000,1300000,['AI 시대, 지금 우리는 어디에 있나',('데이터로 보는 AI와 일의 변화 → AI 역량은 어떻게 측정되는가 → 직접 풀어보기',16,RGBColor(210,215,245),False)],32,WHITE,True)
tb(s,900000,5600000,10400000,400000,'자료: Anthropic Economic Index 연구(2025·2026) · AI 활용평가 유형 분석  |  커리어캐쳐(CareerCatcher)',11,RGBColor(170,180,240))
new.append(s)
# 1
s=base('도입 ① · 지금 우리는 어떤 시대에 있나','AI는 이미 ‘일의 과제’ 단위로 들어와 있습니다','I-1',SRC1)
card(s,X[0],1650000,W,2900000,'약 400만 건의 대화 분석',['Claude.ai 대화를 미국 노동부 O*NET 직업·과제 분류에 연결','“AI가 실제로 어떤 일에 쓰이는가”를 처음 대규모로 측정'])
card(s,X[1],1650000,W,2900000,'소프트웨어·글쓰기 ≈ 절반',['전체 사용의 절반 가까이가 개발·글쓰기 과제','신체 작업 직군(건설·마취 등)은 사용이 거의 없음'])
card(s,X[2],1650000,W,2900000,'직업의 약 36%',['과제의 25% 이상에 AI를 쓰는 직업 비율','과제의 75% 이상 사용 직업은 약 4%'],BLUE)
tb(s,609600,4800000,10858500,700000,'핵심  직업이 통째로 사라지기보다, 직업 안의 ‘과제’부터 AI와 나눠 하기 시작했습니다.',17,BLUE,True)
note(s,'2025년 2월 Anthropic Economic Index 첫 보고서. 개인정보 보호 분석 시스템(Clio)으로 대화를 집계. 단일 플랫폼(Claude.ai) 데이터라는 한계가 있음.')
new.append(s)
# 2
s=base('도입 ② · AI를 쓰는 방식','‘대신 해주는 AI’와 ‘함께 하는 AI’','I-2',SRC1+'\n'+SRC2)
card(s,X[0],1650000,W,2900000,'증강 57% · 자동화 43%',['증강: 배우고, 다듬고, 함께 반복','자동화: 사람 개입 최소로 결과물 수행','(2025.2 보고서 기준)'])
card(s,X[1],1650000,W,2900000,'역량은 아직 다 쓰이지 않음',['Computer·Math 과제의 94%가 이론상 AI 가능','실제 관찰된 사용(커버리지)은 33%','→ 활용 격차가 큼'])
card(s,X[2],1650000,W,2900000,'가장 노출된 직업',['프로그래머 75%','고객상담·데이터 입력 직군 상위','노출 0% 근로자도 약 30%'],BLUE)
tb(s,609600,4800000,10858500,700000,'핵심  AI가 ‘할 수 있는 일’과 실제 ‘하고 있는 일’ 사이의 간격 — 이 간격을 채우는 사람이 경쟁력입니다.',17,BLUE,True)
note(s,'관찰 노출(observed exposure): 이론적 LLM 가능성(Eloundou 외 2023) + 실제 Claude 사용 데이터를 결합, 업무용·자동화 사용에 가중치. 수치는 2026.3 보고서.')
new.append(s)
# 3
s=base('도입 ③ · 일자리에 미치는 영향','일자리 충격은 아직 작지만, 신입 채용부터 신호가 보입니다','I-3',SRC2)
card(s,X[0],1650000,W,2900000,'실업률: 뚜렷한 증가 없음',['2022년 말 이후 고노출 직군의 실업률이 체계적으로 늘지는 않음'])
card(s,X[1],1650000,W,2900000,'22~25세 신규 취업 ↓',['고노출 직군 입직률이 2022년 대비 약 14% 하락','“시사적 근거” 수준 — 확정 아님'],BLUE)
card(s,X[2],1650000,W,2900000,'성장 전망 둔화',['노출 10%p ↑ → 2034년까지 BLS 고용 성장 전망 0.6%p ↓','고노출 근로자: 고학력·고임금·여성 비율 높음'])
tb(s,609600,4800000,10858500,700000,'핵심  ‘AI를 쓸 줄 아는가’가 첫 일자리의 문턱이 되고 있습니다 → 그래서 AI 역량을 측정합니다.',17,BLUE,True)
note(s,'미국 데이터(CPS, BLS) 기반. 보고서도 인과 해석에 신중함을 명시. 한국 상황에 그대로 적용하지 않도록 안내.')
new.append(s)
# 4
SRC3='출처: AI 활용평가(AI Assessment) 미리보기 화면 — 강사 제공 캡처 5종 재구성. 실제 시험 구성·배점은 시행 기관 공지 기준으로 확인 필요'
s=base('도입 ④ · AI 역량은 이렇게 측정됩니다','AI 활용평가: 아는가 → 고르는가 → 해내는가','I-4',SRC3)
rows=[('1. 개념 이해','텍스트 생성형(다음 단어 확률 예측)·이미지 생성형 AI의 기능 차이, 모달리티 구분'),
      ('2. 프롬프트 설계','상황 지문을 읽고 페르소나·기준·출력 형식 등 다중 조건을 모두 갖춘 프롬프트 고르기'),
      ('3. 결과 검증','AI 산출물의 오류(예: 결측치를 0으로 일괄 대체)를 찾아 바로잡기'),
      ('4. 윤리·규제','개인정보·동의 규정을 지키면서 마감을 맞추는 방법(예: 합성 데이터) 판단'),
      ('5. 실습형 과제','첨부 파일(PDF·이미지)을 AI 어시스턴트(질문 횟수 제한)와 함께 분석해 정답 도출')]
y=1600000
for h,b in rows:
    tb(s,590550,y,2600000,700000,h,14,WHITE,True,fill=NAVY); tb(s,3190550,y,8400000,700000,b,13,INK,fill=LIGHT); y+=780000
tb(s,590550,5580000,11010900,520000,'추천 연습 사이트  프로그래머스(programmers.co.kr) — 코딩·AI 관련 문제를 직접 풀며 실전 감각을 기르기 좋습니다.',13,BLUE,True,line=BLUE)
note(s,'형식: 지문(상황) → 4지선다 객관식 + 파일 첨부 실습. 실습형은 AI가 “특정할 수 없다”고 답해도 사람이 조건(예: 중복 주문 시 마지막만 유효)을 다시 적용해 검증해야 함.')
new.append(s)
# 5
def exam(s,num,kind,passage,question,opts,y0=1450000):
    tb(s,590550,y0,4000000,380000,[(f'{num}. {kind}',15,INK,True)],15)
    tb(s,4200000,y0+40000,1500000,300000,'다시 풀 문제 ⚑',9,GRAY,line=GRAY)
    tb(s,590550,y0+420000,11010900,320000,'다음 지문을 읽고 물음에 답하세요.',13,INK)
    bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Emu(640000),Emu(y0+820000),Emu(60000),Emu(900000)); bar.fill.solid(); bar.fill.fore_color.rgb=NAVY; bar.line.fill.background()
    tb(s,780000,y0+800000,10800000,950000,passage,12,GRAY)
    tb(s,590550,y0+1800000,11010900,620000,question,13,INK,True)
    tb(s,590550,y0+2420000,4000000,280000,'답안을 선택하세요 ↴',9,GRAY)
    tb(s,590550,y0+2700000,11010900,1900000,opts,12.5,INK)
s=base('도입 ⑤ · 직접 풀어보기 1','[연습 1] 실제 시험처럼 풀어보기 — 프롬프트 설계','I-5','연습문제: 커리어캐쳐 자체 제작(AI 활용평가 객관식 유형·화면 구성 참고). 정답·해설은 슬라이드 노트')
exam(s,1,'객관식','고용센터 취업지원팀의 김 주임은 참여자 만족도 설문 결과를 정리하기 위해 대화형 AI 도구를 사용했습니다. 그런데 AI의 1차 산출물을 검토한 결과, 응답하지 않은 문항(무응답)을 인공지능이 임의로 ‘매우 불만족(1점)’으로 일괄 처리하여 프로그램 만족도가 실제보다 비정상적으로 낮게 계산되는 오류가 발생했습니다.',
 '위 지문의 오류를 해결하기 위해 1) 설문 분석 전문가 페르소나 부여, 2) 올바른 무응답 처리 기준 제시, 3) 최종 결과의 표 형태 출력이라는 다중 제약 조건을 모두 만족하는 최적의 프롬프트는 무엇인가요?',
 ['①  너는 설문 분석 전문가야. 통계 왜곡을 막기 위해 무응답을 모두 1점으로 유지하되, 문항별 평균 만족도를 한눈에 볼 수 있도록 표 형태로 정리해 줘.',
  '②  너는 취업지원팀 팀장이야. 통계 왜곡을 막기 위해 무응답은 평균 계산에서 제외하고 문항별 응답 수를 함께 적어, 최종 결과를 표 형태로 다시 제시해 줘.',
  '③  너는 설문 분석 전문가야. 통계 왜곡을 막기 위해 무응답은 평균 계산에서 제외하고 문항별 응답 수를 함께 적어, 최종 결과를 표 형태로 다시 제시해 줘.',
  '④  너는 설문 분석 전문가야. 통계 왜곡을 막기 위해 무응답은 평균 계산에서 제외하고, 무응답이 발생하게 된 원인을 서술형 문장으로 자세히 설명해 줘.'])
note(s,'정답 ③. ① 잘못된 처리 기준 유지, ② 페르소나 조건 불일치, ④ 표 출력 조건 누락. 실제 시험도 보기 4개가 한 조건씩만 어긋나도록 출제됨 → 조건을 하나씩 대조하는 습관이 핵심.')
new.append(s)
# 6 (파일 실습형 화면 재현)
s=base('도입 ⑥ · 직접 풀어보기 2','[연습 2] 실제 시험처럼 풀어보기 — 파일 실습형','I-6','연습문제: 커리어캐쳐 자체 제작(AI 활용평가 실습형 화면 구성 참고). 데이터는 가상이며 정답·해설은 슬라이드 노트')
tb(s,590550,1400000,11010900,560000,[('주의사항',12,INK,True),'• 한 명당 하나의 반만 신청 가능합니다. 한 사람이 여러 번 신청한 경우, 가장 마지막 신청 이외에는 모두 취소됩니다. 예) 민준:면접반 → 민준:자소서반 순서로 신청 시 앞의 신청(면접반)은 취소됨'],11,INK)
tb(s,590550,2050000,2300000,1500000,[('파일 목록',11,INK,True),'특강신청.pdf      412','참여자명단.png   236'],11,INK,fill=LIGHT)
tb(s,2950000,2050000,4300000,3300000,[('특강신청.pdf',11,INK,True),'09:02  민준 : 면접반','09:10  서연 : 자소서반','09:15  민준 : 자소서반','09:21  하은 : 면접반','09:33  지우 : 자소서반','09:40  서연 : 면접반','09:52  지우 : 면접반',('참여자명단.png : 민준·서연·하은·지우·도윤 (5명)',10,GRAY,False)],11.5,INK,line=GRAY)
tb(s,7400000,2050000,4200000,3300000,[('AI 어시스턴트  ● 3회 남음',11,INK,True),('나: “특강신청.pdf에서 반별 최종 인원과, 참여자명단.png 중 미신청자를 알려줘”',10.5,NAVY,False),' ',('AI: “면접반 4명, 자소서반 3명입니다. 명단 5명 모두 신청 기록이 있어 미신청자는 없습니다.”',10.5,INK,False)],11,INK,fill=LIGHT)
tb(s,590550,5400000,11010900,800000,[('Q. 주의사항을 적용한 면접반 최종 인원과 미신청자로 옳은 것은?',13,INK,True),'①  면접반 4명 · 미신청자 없음     ②  면접반 3명 · 도윤     ③  면접반 3명 · 미신청자 없음     ④  면접반 2명 · 도윤'],12,INK)
note(s,'정답 ②. 최종: 민준→자소서반, 서연→면접반, 하은→면접반, 지우→면접반 → 면접반 3명, 자소서반 1명. 도윤은 기록 없음. AI는 신청 건수(4·3)를 셌고 명단 대조도 틀림. 실제 시험처럼 AI 질문 횟수가 제한되므로, 규칙을 담아 한 번에 정확히 묻고 결과는 직접 검증해야 함.')
new.append(s)
# 7 해설1
def ans_table(s,rows,y=2250000):
    tb(s,590550,y,1300000,420000,'보기',13,WHITE,True,fill=NAVY); tb(s,1890550,y,1500000,420000,'판정',13,WHITE,True,fill=NAVY); tb(s,3390550,y,8200000,420000,'이유',13,WHITE,True,fill=NAVY)
    y+=440000
    for a,b,c in rows:
        ok=b.startswith('정답')
        tb(s,590550,y,1300000,520000,a,13,INK,True,fill=LIGHT); tb(s,1890550,y,1500000,520000,b,13,BLUE if ok else GRAY,True,fill=LIGHT); tb(s,3390550,y,8200000,520000,c,12.5,INK,fill=LIGHT); y+=560000
    return y
s=base('도입 ⑦ · 정답 해설 1','[연습 1 해설] 정답 ③ — 조건 세 가지를 모두 갖춘 보기','I-7','연습문제: 커리어캐쳐 자체 제작(AI 활용평가 객관식 유형 참고)')
tb(s,590550,1500000,11010900,600000,'풀이 방법  보기마다 조건 ① 페르소나 ② 무응답 처리 기준 ③ 표 출력을 하나씩 대조해, 셋 다 맞는 보기를 고릅니다.',14,INK,fill=LIGHT)
y=ans_table(s,[('①','오답','무응답을 1점으로 유지 → 왜곡을 일으킨 기준을 그대로 둠 (조건 ② 위반)'),('②','오답','“취업지원팀 팀장” 역할 → 설문 분석 전문가 페르소나가 아님 (조건 ① 위반)'),('③','정답','전문가 역할 + 무응답 제외·응답 수 표기 + 표 출력 → 세 조건 모두 충족'),('④','오답','결과 대신 원인을 서술형 문장으로 요청 → 표 출력이 빠짐 (조건 ③ 위반)')])
tb(s,609600,y+150000,10858500,500000,'실무 팁  프롬프트를 보내기 전 “역할·기준·형식” 세 가지를 체크리스트로 확인하세요.',15,BLUE,True)
new.append(s)
# 8 해설2
s=base('도입 ⑧ · 정답 해설 2','[연습 2 해설] 정답 ② — 면접반 3명 · 미신청자 도윤','I-8','연습문제: 커리어캐쳐 자체 제작(AI 활용평가 실습형 유형 참고). 데이터는 가상')
tb(s,590550,1500000,5400000,3500000,[('1단계  사람별 마지막 신청만 남기기',14,NAVY,True),'민준  면접반 → 자소서반  ⇒ 자소서반','서연  자소서반 → 면접반  ⇒ 면접반','하은  면접반  ⇒ 면접반','지우  자소서반 → 면접반  ⇒ 면접반',('최종  면접반 3명 · 자소서반 1명',14,BLUE,True),' ',('2단계  명단과 대조',14,NAVY,True),'명단 5명 중 신청 기록이 없는 사람  ⇒ 도윤'],13,INK,fill=LIGHT)
tb(s,6200000,1500000,5400000,3500000,[('AI 답이 틀린 이유',14,NAVY,True),'• “면접반 4·자소서반 3”은 취소된 신청까지 센 건수','• 명단 대조도 하지 않고 “미신청자 없음”으로 답함',' ',('이렇게 물었다면',14,NAVY,True),'“한 사람이 여러 번 신청했으면 마지막 신청만 남기고, 사람별 최종 반을 표로 보여줘. 그다음 명단과 대조해 기록이 없는 사람을 찾아줘.”'],13,INK,line=NAVY)
tb(s,609600,5250000,10858500,500000,'실무 팁  질문 횟수가 제한될수록 규칙을 질문에 넣고, AI의 답은 중간 과정을 보며 직접 검증하세요.',15,BLUE,True)
new.append(s)
# move after slide index 1 (강사소개)
lst=p.slides._sldIdLst; ids=list(lst)
for i,s in enumerate(new):
    el=ids[len(ids)-len(new)+i]; lst.remove(el); lst.insert(2+i,el)
p.save(out)
